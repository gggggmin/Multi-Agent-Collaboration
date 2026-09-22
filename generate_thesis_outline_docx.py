from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Inches


ROOT = Path(__file__).resolve().parent
OUT_DOCX = ROOT / "毕业论文每章内容详细大纲_强化学习策略说明.docx"


def set_run_font(run, size=12, bold=False):
    run.font.name = "宋体"
    run.font.size = Pt(size)
    run.bold = bold


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    if level == 1:
        p.style = "Heading 1"
        size = 16
    elif level == 2:
        p.style = "Heading 2"
        size = 14
    else:
        p.style = "Heading 3"
        size = 12
    run = p.add_run(text)
    set_run_font(run, size=size, bold=True)
    return p


def add_para(doc, text, size=12):
    p = doc.add_paragraph()
    for line_i, line in enumerate(text.split("\n")):
        if line_i:
            p.add_run().add_break()
        run = p.add_run(line)
        set_run_font(run, size=size)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(item)
        set_run_font(run)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        run = p.add_run(item)
        set_run_font(run)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
        for p in hdr[i].paragraphs:
            for r in p.runs:
                set_run_font(r, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
            for p in cells[i].paragraphs:
                for r in p.runs:
                    set_run_font(r, size=10.5)
    doc.add_paragraph()
    return table


def add_code(doc, code):
    p = doc.add_paragraph()
    run = p.add_run(code)
    run.font.name = "Consolas"
    run.font.size = Pt(10.5)


def build_doc():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("毕业论文每章内容详细大纲\n及强化学习策略说明")
    set_run_font(run, size=20, bold=True)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = sub.add_run("论文题目：基于多智能体协同的需求驱动自动化测试生成与自愈方法研究")
    set_run_font(run, size=12)

    add_heading(doc, "一、论文总定位", 1)
    add_para(doc, "一句话讲清楚：本文研究如何利用 RAG 检索增强、多智能体协作和强化学习策略，自动从需求文档生成高质量测试用例与自动化测试脚本，并通过执行反馈实现测试脚本自愈。")
    add_para(doc, "核心主线不是“我用了大模型生成测试”，而是：用多智能体解决 LLM 测试生成不稳定和幻觉问题，用强化学习解决多轮生成过程中的动态决策问题。")

    add_heading(doc, "二、推荐论文题目", 1)
    add_para(doc, "基于多智能体协同的需求驱动自动化测试生成与自愈方法研究")

    add_heading(doc, "三、章节详细大纲", 1)

    add_heading(doc, "第 1 章 绪论", 2)
    add_para(doc, "本章回答：为什么要做这个题？现在有什么问题？本文解决什么？")
    add_heading(doc, "1.1 研究背景与意义", 3)
    add_para(doc, "软件测试是保障软件质量的重要手段，但传统测试用例设计和自动化脚本编写高度依赖人工经验，存在成本高、效率低、维护困难等问题。随着大语言模型的发展，LLM 已经具备较强的自然语言理解、代码生成和测试设计能力，可以辅助完成需求分析、测试用例生成和自动化脚本生成。")
    add_para(doc, "但是，直接使用 LLM 生成测试用例存在明显问题：")
    add_bullets(doc, [
        "容易遗漏需求中的关键业务规则。",
        "容易生成需求中不存在的功能，也就是幻觉。",
        "生成结果质量不稳定。",
        "自动化脚本一旦页面变化，容易执行失败。",
        "多轮迭代通常靠固定轮次控制，缺少动态优化能力。",
    ])
    add_para(doc, "因此，本文提出一个融合 RAG、多智能体协作、强化学习和执行反馈自愈的自动化测试生成框架。")

    add_heading(doc, "1.2 国内外研究现状", 3)
    add_bullets(doc, [
        "LLM 在测试生成中的研究：大语言模型可以根据需求、代码、接口文档生成测试用例或测试脚本，但仍存在幻觉、覆盖不足和稳定性差的问题。",
        "RAG 在测试生成中的研究：RAG 通过检索需求文档、代码片段和接口说明补充上下文，但单独使用 RAG 不能解决生成结果审计和流程决策问题。",
        "多智能体在软件工程中的研究：多智能体可将复杂任务拆成生成者、审计者、修复者和协调者等角色，更适合处理测试生成这种需要反复检查和修正的任务。",
        "强化学习在软件测试中的研究：强化学习可以根据环境反馈学习策略，适合用于是否继续生成、是否停止迭代、是否切换策略以及失败后选择哪种修复方式等动态决策问题。",
    ])
    add_heading(doc, "1.3 现有研究不足", 3)
    add_numbered(doc, [
        "LLM 生成测试用例容易幻觉。",
        "RAG 只能增强上下文，不能保证生成质量。",
        "多智能体流程常常是固定规则，缺少学习能力。",
        "自动化脚本失败后的修复策略比较单一。",
    ])
    add_heading(doc, "1.4 本文主要研究内容与创新点", 3)
    add_numbered(doc, [
        "构建基于 RAG 的需求检索增强模块。",
        "设计 Creator-Auditor 多智能体测试生成机制。",
        "提出基于 PPO 的自适应生成迭代决策策略 RL-AGS。",
        "构建基于执行反馈的自动化测试脚本自愈机制。",
    ])

    add_heading(doc, "第 2 章 相关理论与技术", 2)
    add_para(doc, "本章回答：本文方法依赖哪些技术？")
    add_heading(doc, "2.1 大语言模型", 3)
    add_para(doc, "介绍 LLM 的自然语言理解、测试用例生成、代码生成和错误分析能力，同时指出 LLM 并不可靠，可能胡编页面元素、接口和断言。论文里不能只吹大模型，必须把局限性讲清楚。")
    add_heading(doc, "2.2 检索增强生成 RAG", 3)
    add_para(doc, "RAG 的基本流程如下：")
    add_code(doc, "需求文档 -> 文档解析 -> 语义切片 -> 向量化存储 -> Top-K 检索 -> 增强 LLM 输入")
    add_para(doc, "本文使用 DocumentParser、SemanticChunker、ChromaDB、sentence-transformers 和 Top-K 检索。RAG 的作用是让 LLM 基于真实需求片段生成测试用例，而不是凭空想。")
    add_heading(doc, "2.3 多智能体系统", 3)
    add_table(doc, ["Agent", "作用"], [
        ("Creator Agent", "根据需求和检索上下文生成测试用例"),
        ("Auditor Agent", "审计测试用例质量，识别遗漏、幻觉和断言问题"),
        ("Page Agent", "生成 Page Object 页面对象"),
        ("Test Agent", "生成 Pytest / Playwright 测试脚本"),
        ("Healing Agent", "修复执行失败的测试脚本"),
        ("Coordinator", "控制整体协作流程"),
    ])
    add_heading(doc, "2.4 强化学习基础", 3)
    add_para(doc, "本文把多智能体测试生成过程看成一个强化学习环境，让智能体根据当前生成质量和审计反馈决定下一步动作。")
    add_table(doc, ["强化学习元素", "在本文中的含义"], [
        ("State 状态", "当前测试生成过程的质量状态"),
        ("Action 动作", "系统下一步要执行的操作"),
        ("Reward 奖励", "当前动作带来的收益或惩罚"),
        ("Policy 策略", "在某个状态下选择动作的规则"),
        ("Environment 环境", "测试生成、审计、执行和反馈流程"),
    ])
    add_para(doc, "本文主强化学习策略采用 PPO，即 Proximal Policy Optimization。PPO 用于学习测试生成过程中的迭代控制策略，核心目标是在保证测试用例质量的前提下减少无效迭代和模型调用成本。")
    add_heading(doc, "2.5 自动化测试技术", 3)
    add_para(doc, "介绍 Pytest、Playwright、Page Object Model、Web 自动化测试、测试脚本执行与错误反馈。本文生成的是 Pytest + Playwright 自动化测试脚本。")

    add_heading(doc, "第 3 章 系统设计与核心方法", 2)
    add_para(doc, "这是论文最重要的一章，回答：系统怎么设计？每个模块怎么工作？强化学习策略具体怎么用？")
    add_heading(doc, "3.1 总体架构", 3)
    add_code(doc, "输入层 -> RAG 知识增强层 -> 多智能体测试生成层 -> 强化学习决策层 -> 脚本生成与执行自愈层")
    add_bullets(doc, [
        "输入层：需求文档、页面信息、接口信息、历史测试结果。",
        "RAG 知识增强层：解析需求、语义切片、向量检索、返回相关需求片段。",
        "多智能体测试生成层：Creator 生成测试用例，Auditor 审计测试用例，Coordinator 组织多轮反馈。",
        "强化学习决策层：判断是否继续迭代、接受当前结果或切换生成策略。",
        "执行自愈层：生成脚本、执行测试、捕获错误并自动修复。",
    ])
    add_heading(doc, "3.2 RAG 需求检索模块", 3)
    add_para(doc, "将需求文档拆成结构化内容，例如功能模块、操作路径、预期结果和异常情况。语义切片尽量按功能模块划分，如登录、搜索、购物车和结算。向量检索使用 embedding 模型将文本转化为向量并存入 ChromaDB，在生成测试用例时召回最相关的需求片段。")
    add_heading(doc, "3.3 Creator-Auditor 多智能体测试生成机制", 3)
    add_para(doc, "Creator 的任务是生成测试用例，输入包括当前测试目标、RAG 检索结果和历史审计反馈，输出包括用例编号、测试标题、前置条件、测试步骤、测试数据、预期结果和用例类型。")
    add_table(doc, ["审计维度", "检查内容"], [
        ("需求一致性", "是否生成了需求中没有的功能"),
        ("需求覆盖率", "是否覆盖主要业务规则"),
        ("步骤完整性", "测试步骤是否可执行"),
        ("断言合理性", "预期结果是否明确"),
        ("边界覆盖", "是否包含异常和边界情况"),
    ])
    add_code(doc, "RAG 检索需求片段 -> Creator 生成测试用例 -> Auditor 审计 -> 反馈问题 -> Creator 重写 -> 下一轮")

    add_heading(doc, "3.4 基于 PPO 的自适应博弈终止策略 RL-AGS", 3)
    add_para(doc, "RL-AGS 全称为 Reinforcement Learning-based Adaptive Game Stopping，即基于强化学习的自适应博弈终止策略。它解决的问题是：Creator-Auditor 到底迭代几轮最合适？")
    add_para(doc, "传统固定轮次策略的问题是：简单需求可能 1 轮就够，跑 3 轮浪费；复杂需求可能 3 轮不够，提前停止质量差；固定规则没有学习能力。本文使用 PPO 根据当前状态动态决定下一步动作。")
    add_heading(doc, "3.4.1 状态空间 State", 3)
    add_code(doc, "s_t = [当前迭代轮次, 最大迭代轮次, Auditor 问题数量, 严重问题数量, 当前需求覆盖率, 当前幻觉率, 上一轮覆盖率提升, 当前测试脚本通过率, 累计 LLM 调用成本]")
    add_table(doc, ["状态变量", "含义"], [
        ("当前轮次", "已经迭代了几次"),
        ("问题数量", "Auditor 发现多少问题"),
        ("严重问题数量", "是否存在关键错误"),
        ("需求覆盖率", "当前用例覆盖了多少需求点"),
        ("幻觉率", "是否生成了需求外内容"),
        ("覆盖率提升", "本轮比上一轮好多少"),
        ("脚本通过率", "生成脚本能否执行"),
        ("调用成本", "已经调用 LLM 多少次或消耗多少 token"),
    ])
    add_heading(doc, "3.4.2 动作空间 Action", 3)
    add_table(doc, ["动作", "含义"], [
        ("a=0：继续当前策略迭代", "当前结果还不够好，让 Creator 根据反馈继续改"),
        ("a=1：接受当前测试用例，停止生成", "当前质量已经够了，避免浪费调用成本"),
        ("a=2：切换生成策略后重新生成", "当前 Prompt 或生成方式效果不好，换策略重新生成"),
    ])
    add_heading(doc, "3.4.3 奖励函数 Reward", 3)
    add_code(doc, "R = α × 覆盖率提升 + β × 审计通过奖励 + γ × 脚本通过率提升 - λ × 幻觉惩罚 - μ × 严重问题惩罚 - η × 调用成本")
    add_para(doc, "奖励函数的人话解释：覆盖率提高、Auditor 通过、脚本能跑通就加分；出现幻觉、严重问题多、LLM 调用太多、迭代很多但质量没提升就扣分。")
    add_heading(doc, "3.4.4 PPO 策略训练", 3)
    add_para(doc, "本文使用 Stable-Baselines3 实现 PPO 算法，将测试生成环境封装为 Gymnasium 环境。PPO 智能体通过多轮实验学习不同状态下的最优动作选择策略。训练目标是在保证测试用例质量的前提下，减少无效迭代和模型调用成本。")
    add_heading(doc, "3.4.5 与固定轮次策略对比", 3)
    add_table(doc, ["方法", "平均迭代轮次", "需求覆盖率", "幻觉率", "调用成本"], [
        ("固定 1 轮", "低", "可能低", "可能高", "低"),
        ("固定 3 轮", "中", "中高", "中", "中"),
        ("固定 5 轮", "高", "不一定更高", "不一定更低", "高"),
        ("RL-AGS", "动态", "高", "低", "更优"),
    ])

    add_heading(doc, "3.5 测试脚本生成模块", 3)
    add_para(doc, "Page Agent 负责生成 Page Object 页面对象，Test Agent 负责生成 Pytest 测试脚本。使用 Page Object Model 可以降低脚本重复度，提高可维护性。")
    add_heading(doc, "3.6 基于执行反馈的自动化脚本自愈机制", 3)
    add_code(doc, "执行测试脚本 -> 捕获错误日志 -> 错误诊断分类 -> 选择修复策略 -> 修改脚本 -> 重新执行验证")
    add_table(doc, ["错误类型", "示例"], [
        ("LocatorError", "找不到按钮或输入框"),
        ("TimeoutError", "页面加载超时"),
        ("AssertionError", "断言失败"),
        ("DataError", "测试数据不合法"),
        ("EnvironmentError", "服务未启动或依赖缺失"),
    ])
    add_table(doc, ["修复策略", "作用"], [
        ("规则修复", "替换常见定位器"),
        ("等待修复", "增加等待或调整超时时间"),
        ("断言修复", "根据页面实际结果重写断言"),
        ("LLM 重写", "调用 LLM 根据错误日志修复"),
        ("回退策略", "回到上一个可执行版本"),
    ])
    add_para(doc, "可选辅助策略 RL-HSS：对于脚本自愈策略选择，可采用 Multi-Armed Bandit / UCB。UCB 适合这种离散修复策略选择场景，因为它会优先选择历史效果好，同时还没有充分尝试过的策略。")

    add_heading(doc, "第 4 章 实验设计与结果分析", 2)
    add_heading(doc, "4.1 实验环境", 3)
    add_bullets(doc, ["Python", "Pytest", "Playwright", "Flask 电商系统", "ChromaDB", "sentence-transformers", "Stable-Baselines3", "Gymnasium"])
    add_para(doc, "被测系统覆盖登录、商品搜索、购物车和结算等典型电商业务模块。")
    add_heading(doc, "4.2 评价指标", 3)
    add_table(doc, ["指标", "含义"], [
        ("需求覆盖率", "生成测试用例覆盖了多少需求点"),
        ("幻觉率", "有多少用例包含需求外内容"),
        ("测试脚本通过率", "生成脚本实际执行成功比例"),
        ("自愈成功率", "执行失败脚本被修复成功比例"),
        ("平均迭代轮次", "每个任务平均迭代多少轮"),
        ("LLM 调用成本", "平均调用次数或 token 成本"),
        ("审计通过率", "Auditor 最终通过比例"),
    ])
    add_heading(doc, "4.3 对比实验", 3)
    add_table(doc, ["方法", "描述"], [
        ("LLM Only", "直接用 LLM 生成"),
        ("LLM + RAG", "加入需求检索"),
        ("Creator-Auditor", "加入多智能体审计"),
        ("Creator-Auditor + RL-AGS", "加入 PPO 动态决策"),
        ("Full Framework", "RAG + 多智能体 + RL + 自愈"),
    ])
    add_heading(doc, "4.4 消融实验", 3)
    add_table(doc, ["消融实验", "对比"], [
        ("RAG 消融", "有 RAG vs 无 RAG"),
        ("Auditor 消融", "有 Auditor vs 无 Auditor"),
        ("RL 消融", "固定轮次 vs RL-AGS"),
        ("自愈消融", "有自愈 vs 无自愈"),
        ("切片策略消融", "语义切片 vs 固定长度切片"),
    ])
    add_heading(doc, "4.5 RL-AGS 实验分析", 3)
    add_para(doc, "重点展示奖励收敛曲线、平均迭代轮次对比、覆盖率对比和调用成本对比。需要证明 PPO 学到的策略确实有用，而不是单纯把轮次堆高。理想结论是：RL-AGS 能够在简单任务中提前停止，在复杂任务中增加迭代轮次，从而减少无效调用并保持较高覆盖率。")
    add_heading(doc, "4.6 自愈实验分析", 3)
    add_para(doc, "分析不同错误类型下的修复效果。一般来说，定位器错误和超时错误更容易修复，复杂业务逻辑错误和需求偏离类错误修复难度更高。")
    add_heading(doc, "4.7 失败案例分析", 3)
    add_bullets(doc, [
        "当需求本身模糊时，RAG 也无法提供明确依据。",
        "当页面缺少稳定定位属性时，自愈难度较高。",
        "当 LLM 生成目标本身偏离需求时，后续脚本修复意义有限。",
        "强化学习训练需要足够样本，冷启动阶段策略不稳定。",
    ])

    add_heading(doc, "第 5 章 总结与展望", 2)
    add_heading(doc, "5.1 工作总结", 3)
    add_numbered(doc, [
        "构建了 RAG 需求增强模块。",
        "设计了 Creator-Auditor 多智能体测试生成机制。",
        "提出了基于 PPO 的 RL-AGS 动态迭代策略。",
        "实现了测试脚本生成与执行反馈自愈。",
        "通过实验验证了方法有效性。",
    ])
    add_heading(doc, "5.2 不足", 3)
    add_bullets(doc, [
        "实验对象主要是电商系统，场景规模有限。",
        "强化学习训练样本数量有限。",
        "LLM 输出仍存在不确定性。",
        "自愈机制对复杂业务逻辑错误修复能力有限。",
        "多智能体调用成本仍然较高。",
    ])
    add_heading(doc, "5.3 展望", 3)
    add_bullets(doc, [
        "扩展到更多类型系统，如后台管理、接口测试和移动端测试。",
        "引入多智能体强化学习 MARL。",
        "优化测试用例优先级排序。",
        "引入更细粒度的代码覆盖率反馈。",
        "构建长期记忆，让系统复用历史修复经验。",
    ])

    add_heading(doc, "四、强化学习策略总览", 1)
    add_table(doc, ["模块", "强化学习策略", "作用", "论文定位"], [
        ("RL-AGS", "PPO", "控制 Creator-Auditor 是否继续迭代、停止或切换策略", "主创新，必须重点写和重点实验"),
        ("RL-HSS", "Multi-Armed Bandit / UCB", "脚本失败后选择规则修复、等待修复、断言修复、LLM 修复或回退策略", "辅助创新，可作为扩展模块"),
    ])
    add_para(doc, "最稳的写法是：主打 PPO，Bandit 辅助。不要把 PPO、DQN、MARL、Bandit 全塞进主线，否则会像技术菜单，答辩时容易被老师追着问每个算法到底贡献在哪里。")
    add_para(doc, "论文主线应压成一句话：本文研究的是在 LLM 生成测试用例的过程中，如何通过多智能体审计减少幻觉，并通过强化学习动态控制生成迭代，从而在测试质量和调用成本之间取得更优平衡。")

    doc.save(OUT_DOCX)
    return OUT_DOCX


if __name__ == "__main__":
    print(build_doc())
