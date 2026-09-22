from copy import deepcopy
from pathlib import Path
import shutil

from docx import Document
from docx.shared import Pt


ROOT = Path(__file__).resolve().parent
TEMPLATE_DOCX = ROOT / "2025级-伍伦贡联合研究院_研究生学位论文相关文件" / "王洋开题答辩文档（中文）.docx"
OUT_DOCX = ROOT / "开题答辩文档_基于多智能体协同的需求驱动自动化测试生成与自愈方法研究.docx"
OUT_OUTLINE = ROOT / "开题答辩PPT大纲_基于多智能体协同的需求驱动自动化测试生成与自愈方法研究.md"


TITLE = "基于多智能体协同的需求驱动自动化测试生成与自愈方法研究"


def set_paragraph_text(paragraph, text, size=12, bold=False):
    paragraph.clear()
    run = paragraph.add_run(text)
    run.font.name = "宋体"
    run.font.size = Pt(size)
    run.bold = bold


def set_cell_text(cell, text, size=10.5):
    cell.text = ""
    lines = text.split("\n")
    first = True
    for line in lines:
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        run = p.add_run(line)
        run.font.name = "宋体"
        run.font.size = Pt(size)


def fill_cover(doc):
    replacements = {
        "研究题目： 基于深度学习的药物-靶点结合亲和力预测": f"研究题目： {TITLE}",
        "学生姓名：            王洋": "学生姓名：            【待填写】",
        "学号： 2024124258": "学号： 【待填写】",
        "主要的：       计算机科学": "主要的：       计算机科学",
        "导师：陈矛": "导师：【待填写】",
        "日期：              2025年7月6日": "日期：              2026年7月21日",
        "2025年7月": "2026年7月",
    }
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text in replacements:
            set_paragraph_text(paragraph, replacements[text])


CONTENT = {
    "challenges": """1. 需求理解与测试生成之间存在语义断层
传统自动化测试依赖人工将需求文档转化为测试用例和脚本，转换过程成本高、主观性强。大语言模型虽然具备自然语言理解和代码生成能力，但直接生成测试用例时容易遗漏业务规则、边界条件和异常路径，导致测试集与真实需求之间存在偏差。

2. LLM 测试生成存在幻觉与不可控问题
LLM 在测试生成中可能编造需求中不存在的功能、接口、页面元素或断言逻辑。若缺少审计机制，生成结果看似完整，实际却可能无法执行或无法验证核心业务。单次 Prompt 调用本质上是黑盒生成，质量波动大，难以满足研究与工程落地要求。

3. 多轮生成流程固定，缺少动态优化能力
现有 Creator-Auditor 迭代通常采用固定轮次，质量已经足够时仍继续调用模型，造成成本浪费；质量不足时又可能提前停止，导致结果不稳定。这种硬编码流程没有学习能力，不能根据审计反馈、覆盖率变化和调用成本动态权衡。

4. 自动化脚本维护成本高，自愈策略单一
Web 自动化脚本容易因页面结构变化、元素定位器失效、等待时间不足和断言变化而失败。若每次失败都依赖人工修复或直接调用 LLM 重写，成本高且缺少可复用经验。如何根据错误类型选择合适修复策略，是测试生成闭环落地的关键挑战。""",
    "significance": """本研究面向需求驱动的自动化测试生成任务，构建融合 RAG、多智能体协作、强化学习决策和执行反馈自愈的测试生成框架，具有以下意义：

1. 提升 LLM 测试生成的可靠性
通过 RAG 检索需求上下文，降低模型因上下文不足造成的遗漏和幻觉；通过 Creator-Auditor 对抗审计机制，将测试生成从单次输出改造为可检查、可迭代的质量提升过程。

2. 推动测试生成流程从 Prompt 编排走向算法化优化
本文将强化学习引入多智能体测试生成过程，使用 RL-AGS 动态判断继续迭代、接受输出或切换生成策略，使系统能够在覆盖率、质量和调用成本之间进行学习型权衡。这个点是论文能不能站住的硬核主线，不能只停留在“我调了几个 Agent”。

3. 降低自动化测试脚本维护成本
执行反馈自愈模块能够捕获失败信息、诊断错误类型，并结合规则修复、LLM 重写和策略选择机制完成修复与重执行。该机制有助于提升生成脚本在页面变化和运行环境波动下的可维护性。

4. 形成可复现实验框架
研究以电商系统为实验对象，覆盖登录、搜索、购物车、结算等典型业务流程，可围绕需求覆盖率、幻觉率、脚本通过率、自愈成功率、平均迭代轮次和调用成本开展系统评价，为后续扩展到更多业务系统提供基础。""",
    "goals": """研究目标
1. 构建面向需求文档的 RAG 检索增强模块，实现需求解析、语义切片、向量存储与 Top-K 上下文召回，为测试生成提供可信输入。
2. 设计 Creator-Auditor 多智能体协作机制，使 Creator 负责生成正向、逆向和边界测试用例，Auditor 从需求一致性、覆盖完整性、断言合理性和可执行性等维度进行审计。
3. 提出基于强化学习的自适应博弈终止策略 RL-AGS，使用 PPO 对迭代过程进行动态决策，替代固定轮次机制。
4. 设计基于执行反馈的测试脚本自愈机制，结合错误诊断、规则修复、LLM 修复和策略选择，提升脚本执行通过率。
5. 通过对比实验和消融实验验证框架在覆盖率、幻觉抑制、执行通过率和成本控制方面的有效性。

预期成果
1. 一个可运行的自动化测试生成原型系统，包括 RAG、Creator-Auditor、POM 脚本生成、执行自愈和实验评估模块。
2. 一套基于 PPO 的 RL-AGS 决策模型，以及与固定轮次策略的对比实验结果。
3. 一套自愈策略选择机制，输出不同错误类型下的修复成功率和策略有效性分析。
4. 完整实验报告与论文材料，包括系统架构图、流程图、奖励曲线、覆盖率对比图、消融实验表和失败案例分析。""",
    "theory": """本研究采用“需求检索增强—多智能体生成审计—强化学习决策—执行反馈自愈”的总体策略。

1. 检索增强生成理论
需求文档经解析与语义切片后进入向量数据库，生成测试用例前召回相关上下文。RAG 的作用不是装饰门面，而是给 LLM 加业务边界，减少凭空编造。

2. 多智能体协同理论
将测试生成拆分为 Creator、Auditor、Page Agent、Test Agent、Healing Agent 和 Coordinator 等角色。Creator 负责生成测试用例，Auditor 负责一致性与完备性审计，Page Agent 与 Test Agent 负责自动化脚本生成，Healing Agent 根据执行失败信息修复脚本。后续可基于 AutoGen 的 AssistantAgent 与 GroupChat 重构通信层，提高协作流程的标准化和可扩展性。

3. 强化学习决策理论
将 Creator-Auditor 迭代过程建模为马尔可夫决策过程。状态包括当前轮次、审计问题数量、严重问题数量、需求覆盖率、执行通过率和调用成本；动作包括继续迭代、接受输出和切换生成策略；奖励函数综合覆盖率提升、审计通过、幻觉惩罚和调用成本惩罚。采用 Stable-Baselines3 中的 PPO 作为主要算法。

4. 自动化测试与自愈理论
测试脚本采用 Pytest 与 Playwright，遵循 Page Object Model 组织页面对象和测试逻辑。自愈模块根据错误日志识别定位器失效、等待超时、断言失败和环境异常，并选择规则修复、LLM 重写或回退策略完成闭环验证。""",
    "literature": """采用上述方法的原因在于：LLM 直接生成测试用例虽然快，但质量不可控；多智能体审计能够补足单 Agent 的盲区；RAG 能把生成过程约束在真实需求上下文内；强化学习则解决固定迭代策略僵硬的问题。

已阅读和参考的文献主要包括：
[1] 曹鹏, 陈刚, 季学纯, 刘歆一, 温广琪, 杨金柱. 面向测试用例生成的大模型高效微调方法. 计算机应用, 2025, 45(3): 725-731.
评论：说明大模型在测试用例生成中的可行性，但侧重微调，未充分解决生成过程的动态审计和执行反馈问题。

[2] 罗嘉乐, 王涛, 程良伦, 等. 基于大语言模型的组件化工业软件系统测试用例生成方法. 广东工业大学学报, 2025.
评论：采用分步生成和修正思路，与本文 Creator-Auditor 迭代机制相关，但本文进一步引入强化学习决策。

[3] 汪莹, 字千成, 彭鑫, 娄一翎. 基于大语言模型的故障复现测试用例生成方法. 软件学报, 2025.
评论：证明 RAG 对测试生成上下文补充有价值，为本文需求检索增强和错误反馈修复提供参考。

[4] Hallucination to Consensus: Multi-Agent LLMs for End-to-End JUnit Test Generation, arXiv, 2025.
评论：多智能体共识机制可降低测试生成幻觉，支撑本文使用 Auditor 审计生成结果的设计。

[5] From LLMs to LLM-based Agents for Software Engineering: A Survey of Current, Challenges and Future, arXiv, 2024.
评论：系统梳理 LLM Agent 在软件工程任务中的应用，为本文多智能体角色划分提供理论背景。

[6] Retrieval-Augmented Test Generation: How Far Are We?, arXiv, 2024.
评论：表明 RAG 在测试生成中能改善覆盖效果，但单纯 RAG 仍缺少审计、执行反馈和策略优化。

[7] Bylina, B., & Antończak, A. Analysis of End-to-End Test Automation Tools Based on Selenium WebDriver and Playwright. FedCSIS, 2024.
评论：为本文选用 Playwright 作为 Web 自动化执行工具提供依据。

[8] Gahlot, S., Bairi, A. R., & Saminathan, M. Self-Healing Automation with Reinforcement Learning: Adaptive Test Scripts Using PPO and Dynamic XPath in Playwright. JAIGS, 2024.
评论：强化学习用于脚本自愈具有可行性，本文将该思想扩展到生成迭代决策和修复策略选择。""",
    "validation": """验证计划
1. 实验场景：构建电商 Web 系统作为被测对象，覆盖登录、商品搜索、购物车、结算等模块，并设计正常流程、异常流程和边界条件。
2. 对比方法：设置 LLM Only、LLM+RAG、固定轮次 Creator-Auditor、Creator-Auditor+RL、带自愈与不带自愈等对比组。
3. 评价指标：需求覆盖率、幻觉率、测试脚本通过率、自愈成功率、平均迭代轮次、平均生成耗时、LLM 调用成本和审计通过率。
4. 消融实验：分别验证 RAG、Auditor、RL-AGS、语义切片、自愈模块和不同 LLM 模型对整体效果的影响。
5. 结果分析：输出覆盖率柱状图、幻觉率对比表、RL 奖励收敛曲线、平均迭代轮次对比图和失败案例分析。

实施计划
1. 编程语言：Python。
2. 测试框架：Pytest、Playwright、Page Object Model。
3. 多智能体框架：现阶段已有自实现 Agent 原型，后续计划集成 Microsoft AutoGen。
4. RAG 工具：ChromaDB、sentence-transformers/all-MiniLM-L6-v2，必要时降级为 TF-IDF 检索。
5. 强化学习工具：Gymnasium 定义环境，Stable-Baselines3 实现 PPO；自愈策略选择可采用 UCB 或 Thompson Sampling。
6. 实验环境：Windows/Python 环境，结合本地 Flask 电商系统与脚本执行沙箱完成实验。""",
    "novelty": """1. 面向测试生成的 Creator-Auditor 对抗审计机制
本文不是简单让 LLM 一把梭生成测试，而是将生成与审计拆成两个角色，通过一致性检查、完备性检查和反馈重写减少幻觉和遗漏。

2. 基于强化学习的自适应博弈终止策略
将测试生成迭代过程建模为决策问题，使用 PPO 根据审计状态动态选择继续、停止或切换策略。相比固定 3 轮迭代，该方法能够学习质量收益与调用成本之间的平衡，这是本文的核心算法创新。

3. RAG 与多智能体协同的需求约束生成
通过语义切片和向量检索把真实需求片段注入生成过程，使测试用例生成更贴合业务规则，降低凭空生成页面元素、接口或断言的概率。

4. 执行反馈驱动的自动化脚本自愈闭环
系统将脚本执行结果反馈给修复模块，并根据错误类型选择定位器修复、等待策略调整、断言重写或 LLM 修复，实现从用例生成到脚本执行再到修复验证的闭环。

5. 面向工程落地的可复现实验体系
本文以具体电商系统为对象，设计多组对比实验和消融实验，不只讲概念，也能用覆盖率、幻觉率、通过率和成本指标说话。答辩时这比空喊“智能化”有杀伤力得多。""",
}


TIMELINE = [
    ("2026年7月-2026年8月", "完成开题材料、补充核心文献调研，明确 RAG、多智能体、强化学习和自愈模块的研究边界。"),
    ("2026年9月-2026年10月", "完成 AutoGen 通信层改造，完善 Creator-Auditor 协作流程和电商系统需求数据集。"),
    ("2026年11月-2026年12月", "实现 RL-AGS 环境建模、PPO 训练脚本和固定轮次对比实验，输出奖励收敛曲线。"),
    ("2027年1月-2027年2月", "完善执行反馈自愈模块，实现错误分类、规则修复、LLM 修复和策略选择实验。"),
    ("2027年3月-2027年4月", "开展完整对比实验和消融实验，整理覆盖率、幻觉率、通过率、成本和自愈成功率数据。"),
    ("2027年5月", "完成论文初稿，补充图表、实验分析、失败案例和总结展望。"),
]


def fill_tables(doc):
    table0, table1, table2 = doc.tables
    set_cell_text(table0.cell(1, 0), CONTENT["challenges"])
    set_cell_text(table0.cell(3, 0), CONTENT["significance"])
    set_cell_text(table0.cell(5, 0), CONTENT["goals"])

    set_cell_text(table1.cell(1, 0), CONTENT["theory"])
    set_cell_text(table1.cell(3, 0), CONTENT["literature"])
    set_cell_text(table1.cell(5, 0), CONTENT["validation"])
    for idx, (period, work) in enumerate(TIMELINE, start=8):
        set_cell_text(table1.cell(idx, 0), period)
        set_cell_text(table1.cell(idx, 1), work)
    set_cell_text(table1.cell(15, 0), CONTENT["novelty"])

    set_cell_text(table2.cell(1, 0), "【导师填写】")
    set_cell_text(table2.cell(2, 0), "结论")
    set_cell_text(table2.cell(2, 1), "【导师填写】")


PPT_OUTLINE = f"""# 开题答辩 PPT 大纲

参考模板：`王洋开题.pptx`，建议沿用 17 页左右的节奏，不要做成 30 页论文压缩包。

## 1. 题目页
- 题目：{TITLE}
- 学生姓名、学号、专业、指导教师、汇报日期

## 2. 目录
- 研究背景
- 研究现状
- 研究目标
- 技术路线
- 论文大纲

## 3. 研究背景
- 人工编写自动化测试用例成本高、维护压力大
- LLM 具备需求理解、用例生成和代码生成能力
- 单次 LLM 生成存在幻觉、遗漏和不可控问题

## 4. 研究领域定位
- 需求驱动自动化测试生成
- LLM Agent for Software Engineering
- RAG、多智能体协作、强化学习决策与测试自愈的交叉方向

## 5. 问题一：需求上下文不足
- 现象：LLM 容易漏掉业务规则、边界条件、异常路径
- 方案：需求解析、语义切片、向量检索、Top-K 上下文增强

## 6. 问题二：生成结果存在幻觉
- 现象：编造页面元素、接口、断言或不存在的流程
- 方案：Creator 生成，Auditor 从一致性、完备性、可执行性审计

## 7. 研究现状
- LLM 测试生成：效率高，但稳定性不足
- 多智能体测试生成：能引入审计与共识，但常缺少学习型决策
- RAG 测试生成：改善上下文，但不能单独解决流程控制
- 强化学习测试优化：适合动态策略选择，但与测试生成闭环结合仍有空间

## 8. 现有方法局限
- 固定轮次迭代浪费调用成本或提前停止
- Prompt 编排过重，算法贡献不足
- 测试脚本失败后修复策略单一
- 实验常停留在生成质量，缺少执行反馈闭环

## 9. 研究目标
- 构建 RAG 增强的需求驱动测试生成框架
- 设计 Creator-Auditor 多智能体审计机制
- 提出基于 PPO 的 RL-AGS 动态迭代决策
- 构建执行反馈驱动的脚本自愈机制

## 10. 总体架构
- 输入层：需求文档、页面信息、历史执行结果
- 知识增强层：语义切片、ChromaDB、Top-K 检索
- 多智能体层：Creator、Auditor、Page Agent、Test Agent、Healing Agent
- 决策优化层：RL-AGS、RL-HSS
- 执行反馈层：Pytest、Playwright、错误诊断与重执行

## 11. 核心模块一：RAG 检索增强
- DocumentParser：解析需求模块、操作路径、预期结果
- SemanticChunker：语义切片与固定切片对比
- Retriever：向量检索召回生成上下文

## 12. 核心模块二：Creator-Auditor 博弈
- Creator：生成正向、逆向、边界值测试用例
- Auditor：检查需求一致性、覆盖完整性、断言合理性
- Coordinator：组织多轮反馈与重写

## 13. 核心模块三：RL-AGS 动态决策
- State：轮次、问题数、严重问题数、覆盖率、通过率、调用成本
- Action：继续迭代、接受输出、切换策略
- Reward：覆盖率提升奖励、审计通过奖励、幻觉惩罚、成本惩罚
- 算法：Stable-Baselines3 PPO

## 14. 核心模块四：执行自愈
- 错误捕获：定位器失效、等待超时、断言失败、环境异常
- 修复策略：规则替换、增加等待、断言重写、LLM 重写、版本回退
- 可扩展：Multi-Armed Bandit 选择最优修复策略

## 15. 实验设计
- 场景：电商系统登录、搜索、购物车、结算
- 对比：LLM Only、LLM+RAG、固定 Creator-Auditor、Creator-Auditor+RL、带/不带自愈
- 指标：需求覆盖率、幻觉率、脚本通过率、自愈成功率、平均迭代轮次、调用成本

## 16. 论文大纲
- 第1章 绪论
- 第2章 相关理论与技术
- 第3章 系统设计与核心算法
- 第4章 实验与结果分析
- 第5章 总结与展望

## 17. 总结页
- 本研究将测试生成从“单次 Prompt 输出”升级为“检索增强、多智能体审计、强化学习决策、执行反馈自愈”的闭环框架
- 重点贡献：RL-AGS 动态迭代决策、Creator-Auditor 审计机制、执行反馈自愈
- 结束语：恳请各位老师批评指正
"""


def main():
    shutil.copyfile(TEMPLATE_DOCX, OUT_DOCX)
    doc = Document(OUT_DOCX)
    fill_cover(doc)
    fill_tables(doc)
    doc.save(OUT_DOCX)
    OUT_OUTLINE.write_text(PPT_OUTLINE, encoding="utf-8")
    print(OUT_DOCX)
    print(OUT_OUTLINE)


if __name__ == "__main__":
    main()
