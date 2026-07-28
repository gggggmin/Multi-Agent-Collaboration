from pathlib import Path
import shutil

from docx import Document
from docx.shared import Pt


ROOT = Path(__file__).resolve().parent
TEMPLATE = (
    ROOT
    / "2025级-伍伦贡联合研究院_研究生学位论文相关文件"
    / "2024级-伍伦贡联合研究院_研究生学位论文相关文件"
    / "Thesis Proposal for JI-template-2025.docx"
)
OUT = ROOT / "Thesis Proposal for JI_Completed_Multi-Agent_RL_Test_Generation.docx"

TITLE = "Research on an Automated Test Generation Framework Based on Multi-Agent Collaboration and Reinforcement Learning"


def set_para(paragraph, text, size=12, bold=False):
    paragraph.clear()
    run = paragraph.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.bold = bold


def set_cell(cell, text, size=10.5):
    cell.text = ""
    first = True
    for line in text.split("\n"):
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        run = p.add_run(line)
        run.font.name = "Times New Roman"
        run.font.size = Pt(size)


def fill_cover(doc):
    replacements = {
        "Research Title：": f"Research Title: {TITLE}",
        "Student Name：": "Student Name: To be filled",
        "Student Number：": "Student Number: To be filled",
        "Major：": "Major: Computer Science",
        "Supervisor：": "Supervisor: To be filled",
        "Date：": "Date: 21 July 2026",
        "July 2025": "July 2026",
    }
    for p in doc.paragraphs:
        t = p.text.strip()
        if t in replacements:
            set_para(p, replacements[t])


PROBLEMS = """This research focuses on automated test case and test script generation from requirement documents by using large language models (LLMs), retrieval-augmented generation (RAG), multi-agent collaboration, reinforcement learning and execution-feedback-based self-healing. The key research problems and challenges are as follows.

1. Semantic gap between requirements and executable tests
Manual test design requires engineers to interpret business requirements, identify test scenarios, define test data, and implement executable scripts. This process is expensive and highly dependent on human experience. Although LLMs can understand natural language and generate code, they may still miss hidden business rules, boundary conditions and negative scenarios when requirements are long or ambiguous.

2. Hallucination and instability in LLM-generated test cases
LLMs may generate functions, UI elements, workflows or assertions that do not exist in the original requirements. Such hallucinated tests may look reasonable in natural language but fail during execution or verify the wrong behavior. Therefore, a single-pass LLM generation pipeline is not reliable enough for research-grade or engineering-grade test generation.

3. Inefficient fixed-round multi-agent iteration
Existing multi-agent generation-and-review workflows usually rely on a fixed number of iterations. This is a blunt strategy: simple requirements may not need multiple rounds, while complex requirements may still require further refinement after the predefined number of rounds. Fixed iteration cannot learn from audit feedback, coverage improvement, execution success, or LLM invocation cost.

4. High maintenance cost of generated automated test scripts
Generated Web test scripts may fail because of locator changes, timeout issues, assertion mismatch, invalid test data, or environmental errors. If every failure is handled manually or by blindly calling an LLM to rewrite the whole script, the maintenance cost remains high. A more practical framework should diagnose failure types and select appropriate repair strategies based on execution feedback.

5. Lack of systematic evaluation
Many LLM-based test generation studies focus only on generated text quality. This research must evaluate not only whether the test cases look reasonable, but also whether they cover requirements, avoid hallucination, produce executable scripts, reduce unnecessary iterations, and improve self-healing success rate."""


SIGNIFICANCE = """The proposed research is significant from both academic and practical perspectives.

1. Improving the reliability of LLM-based test generation
By combining RAG with a Creator-Auditor multi-agent mechanism, the framework constrains generation using retrieved requirement evidence and checks generated test cases against requirement consistency, coverage, boundary conditions and executability. This can reduce hallucination and requirement omission, which are two core weaknesses of direct LLM generation.

2. Turning prompt-based workflows into decision-optimized workflows
A purely prompt-based pipeline is too weak as a thesis contribution. The proposed RL-AGS strategy models the Creator-Auditor iteration as a reinforcement learning problem and uses PPO to learn when to continue, stop or switch generation strategy. This gives the thesis a clear algorithmic contribution rather than only an engineering integration.

3. Reducing automated test maintenance cost
The execution-feedback-based self-healing module diagnoses script failures and applies repair strategies such as rule-based locator replacement, waiting adjustment, assertion revision, LLM-based rewriting and fallback. This helps generated scripts remain useful when the application interface or runtime environment changes.

4. Providing a reproducible experimental framework
The research uses an e-commerce Web application as the target system, covering login, search, shopping cart and checkout modules. It supports measurable evaluation using requirement coverage, hallucination rate, script pass rate, self-healing success rate, average iteration rounds and LLM invocation cost.

5. Supporting future intelligent software testing research
The proposed framework can be extended to API testing, mobile testing, regression testing, test prioritization and multi-agent reinforcement learning. It provides a practical basis for studying how LLM agents can be made more controllable, auditable and cost-aware."""


OBJECTIVES = """The objectives of this research are:

1. To design a RAG-based requirement retrieval module that parses requirement documents, performs semantic chunking, stores requirement chunks in a vector database, and retrieves relevant contexts for test generation.
2. To build a Creator-Auditor multi-agent test generation mechanism, where the Creator generates positive, negative and boundary test cases, and the Auditor reviews them from requirement consistency, coverage completeness, assertion correctness and executability.
3. To propose RL-AGS, a reinforcement-learning-based adaptive game stopping strategy, which uses PPO to dynamically decide whether to continue iteration, accept the current output or switch the generation strategy.
4. To implement an automated test script generation module based on Page Object Model, Pytest and Playwright.
5. To develop an execution-feedback-based self-healing mechanism for diagnosing and repairing failed generated scripts.
6. To evaluate the proposed framework through comparison experiments and ablation studies on an e-commerce Web application.

The planned outcomes are:

1. A runnable prototype framework integrating requirement retrieval, multi-agent test generation, RL-based iteration control, script generation and self-healing.
2. A PPO-based RL-AGS model with experimental results comparing it against fixed-round iteration strategies.
3. A self-healing strategy selection mechanism with analysis of repair success rate under different failure types.
4. A complete experimental report including architecture diagrams, workflow diagrams, reward curves, coverage comparison charts, ablation tables and failure case analysis.
5. A master thesis that clearly explains the research problem, proposed method, implementation, experiments, results and limitations."""


METHODS = """This research adopts a layered and feedback-driven methodology consisting of five major components.

1. Requirement retrieval and knowledge enhancement
Requirement documents are parsed into structured units such as functional module, operation path, expected result and exception condition. The parsed content is divided into semantic chunks and encoded by an embedding model. ChromaDB is used as the vector database. For each generation task, Top-K relevant requirement chunks are retrieved and injected into the LLM prompt to reduce context omission and hallucination.

2. Creator-Auditor multi-agent collaboration
The test generation process is decomposed into specialized agents. The Creator Agent generates test cases from the retrieved requirement context. The Auditor Agent evaluates generated cases according to requirement consistency, coverage completeness, boundary condition coverage, assertion correctness and execution feasibility. If the audit result is not satisfactory, feedback is sent to the Creator for refinement.

3. RL-AGS: PPO-based adaptive game stopping
The Creator-Auditor iterative process is modeled as a Markov decision process. The state contains the current iteration round, number of audit issues, number of severe issues, requirement coverage, hallucination rate, coverage improvement, script pass rate and accumulated LLM cost. The action space includes continuing the current iteration, accepting the current output, and switching generation strategy. The reward function combines coverage gain, audit pass bonus and script pass improvement, while penalizing hallucination, severe issues and LLM invocation cost. PPO implemented with Stable-Baselines3 is used as the main reinforcement learning algorithm.

4. Automated script generation
After test cases pass audit, the Page Agent generates Page Object Model classes and the Test Agent generates Pytest/Playwright scripts. The POM design reduces code duplication and improves maintainability by separating page operations from test assertions.

5. Execution-feedback-based self-healing
Generated scripts are executed in a controlled testing environment. Failure logs are classified into locator errors, timeout errors, assertion errors, data errors and environment errors. Repair strategies include rule-based locator replacement, waiting adjustment, assertion revision, LLM-based rewriting and fallback. A Multi-Armed Bandit strategy such as UCB can be used as an auxiliary method for selecting the most promising repair strategy based on historical repair success and cost."""


LITERATURE = """The selected methods are adopted because direct LLM generation alone is not reliable enough for automated testing. RAG supplies requirement evidence, multi-agent collaboration introduces independent review, PPO learns dynamic iteration control, and self-healing closes the loop between generation and execution.

The literature reviewed so far includes:

[1] Cao, P., Chen, G., Ji, X., Liu, X., Wen, G., & Yang, J. (2025). Efficient fine-tuning methods of large models for test case generation. Journal of Computer Applications, 45(3), 725-731.
Comment: This work demonstrates the feasibility of LLMs for test case generation, but it mainly focuses on model adaptation and does not sufficiently address dynamic auditing and execution feedback.

[2] Luo, J., Wang, T., Cheng, L., et al. (2025). A large-language-model-based test case generation method for componentized industrial software systems. Journal of Guangdong University of Technology.
Comment: The stepwise generation-revision-code-generation process is related to the proposed Creator-Auditor mechanism. This research further introduces reinforcement learning to optimize the iteration process.

[3] Wang, Y., Zi, Q., Peng, X., & Lou, Y. (2025). A large-language-model-based method for generating failure reproduction test cases. Journal of Software.
Comment: This work supports the value of using RAG and error context for test generation, which is relevant to the requirement retrieval and healing design of this thesis.

[4] Hallucination to Consensus: Multi-Agent LLMs for End-to-End JUnit Test Generation. arXiv preprint, 2025.
Comment: This paper shows that multi-agent consensus can reduce hallucination in test generation, supporting the design of an Auditor role in this research.

[5] From LLMs to LLM-based Agents for Software Engineering: A Survey of Current, Challenges and Future. arXiv preprint, 2024.
Comment: This survey provides a general theoretical background for applying LLM-based agents to software engineering tasks, including testing.

[6] Retrieval-Augmented Test Generation: How Far Are We? arXiv preprint, 2024.
Comment: This study shows that RAG can improve test generation, but RAG alone cannot control iterative generation quality or execution repair.

[7] Bylina, B., & Antonczak, A. (2024). Analysis of end-to-end test automation tools based on the examples of Selenium WebDriver and Playwright. FedCSIS.
Comment: This paper provides support for selecting Playwright as the Web automation execution tool.

[8] Gahlot, S., Bairi, A. R., & Saminathan, M. (2024). Self-healing automation with reinforcement learning: adaptive test scripts using PPO and dynamic XPath in Playwright. Journal of Artificial Intelligence General Science.
Comment: This work indicates that reinforcement learning can be applied to self-healing test automation. This thesis extends the idea to both iteration control and repair strategy selection.

Overall, the reviewed literature confirms that LLM-based test generation, RAG, multi-agent collaboration and reinforcement learning are individually promising. The gap is the lack of a unified framework that combines requirement retrieval, agent-based audit, RL-based iteration control and execution-feedback-based self-healing."""


VALIDATION = """The proposed solution will be validated through implementation, comparison experiments and ablation studies.

1. Experimental target system
An e-commerce Web application will be used as the experimental subject. It contains representative modules such as user login, product search, shopping cart and checkout. These modules provide positive paths, negative paths and boundary conditions suitable for evaluating automated test generation.

2. Baseline methods
The proposed framework will be compared with several baseline methods:
- LLM Only: directly generating test cases using an LLM.
- LLM + RAG: adding requirement retrieval without multi-agent audit.
- Creator-Auditor with fixed rounds: using a fixed number of review iterations.
- Creator-Auditor + RL-AGS: using PPO-based dynamic iteration control.
- Full framework: RAG + multi-agent generation + RL-AGS + script generation + self-healing.

3. Evaluation metrics
The main metrics include requirement coverage, hallucination rate, script pass rate, self-healing success rate, average iteration rounds, generation time, LLM invocation cost and final audit pass rate.

4. Ablation studies
Ablation experiments will be conducted to evaluate the contribution of RAG, Auditor, RL-AGS, semantic chunking and self-healing. Fixed-round strategies such as one round, three rounds and five rounds will be compared with RL-AGS to verify whether PPO can achieve a better balance between quality and cost.

5. Implementation plan
The system will be implemented in Python. Pytest and Playwright will be used for automated Web test execution. The target application is implemented with Flask. ChromaDB and sentence-transformers will be used for vector retrieval. Gymnasium will be used to define the RL environment, and Stable-Baselines3 will be used to train the PPO policy. AutoGen may be used to standardize the multi-agent communication layer.

6. Expected result presentation
The experimental section will include requirement coverage tables, hallucination rate comparison, script pass rate comparison, RL reward convergence curves, average iteration round comparison, LLM cost comparison and failure case analysis."""


NOVELTY = """The novelty of this research lies in the following aspects.

1. A Creator-Auditor mechanism for requirement-consistent test generation
The framework does not simply ask an LLM to generate tests in one pass. Instead, it separates generation and audit into different agents. The Auditor checks requirement consistency, coverage completeness, assertion correctness and executability, which helps reduce hallucinated tests and missing scenarios.

2. PPO-based adaptive iteration control for multi-agent test generation
The proposed RL-AGS strategy models the generation-audit loop as a reinforcement learning decision problem. It uses PPO to dynamically choose whether to continue, stop or switch strategy according to audit feedback, coverage improvement, script pass rate and invocation cost. This is the main algorithmic contribution of the thesis.

3. Integration of RAG, multi-agent collaboration and reinforcement learning
Existing methods often use RAG, agents or RL separately. This research integrates them into one closed-loop framework for requirement-driven automated test generation. RAG provides evidence, agents provide role-based generation and audit, and RL provides adaptive process optimization.

4. Execution-feedback-based self-healing of generated scripts
The framework extends test generation to actual script execution and repair. It diagnoses failure types and applies targeted repair strategies, making the generated scripts more practical than text-only test cases.

5. Cost-aware evaluation
The research does not only evaluate whether generated tests look correct. It also evaluates iteration rounds, LLM invocation cost and self-healing success rate. This is important because an intelligent testing framework must be both effective and economically usable."""


TIMELINE = [
    ("July 2026 - August 2026", "Complete proposal preparation, refine the research scope, and review literature on LLM-based test generation, RAG, multi-agent systems, reinforcement learning and self-healing test automation."),
    ("September 2026 - October 2026", "Improve the RAG module and Creator-Auditor workflow; prepare the e-commerce requirement dataset and target Web application scenarios."),
    ("November 2026 - December 2026", "Implement the RL-AGS Gymnasium environment, define state/action/reward, train the PPO policy, and compare it with fixed-round iteration strategies."),
    ("January 2027 - February 2027", "Complete automated script generation and execution-feedback-based self-healing; implement failure diagnosis and repair strategy selection."),
    ("March 2027 - April 2027", "Conduct full comparison experiments and ablation studies; collect data on coverage, hallucination, pass rate, self-healing success rate and cost."),
    ("May 2027", "Complete the first draft of the thesis, including figures, experimental results, discussion, limitations and future work."),
]


def fill_doc(doc):
    fill_cover(doc)
    table0, table1, table2 = doc.tables
    set_cell(table0.cell(1, 0), PROBLEMS)
    set_cell(table0.cell(3, 0), SIGNIFICANCE)
    set_cell(table0.cell(5, 0), OBJECTIVES)

    set_cell(table1.cell(1, 0), METHODS)
    set_cell(table1.cell(3, 0), LITERATURE)
    set_cell(table1.cell(5, 0), VALIDATION)
    for row_idx, (period, content) in enumerate(TIMELINE, start=8):
        set_cell(table1.cell(row_idx, 0), period)
        set_cell(table1.cell(row_idx, 1), content)
    set_cell(table1.cell(15, 0), NOVELTY)

    set_cell(table2.cell(1, 0), "To be completed by the supervisor.")
    set_cell(table2.cell(2, 0), "Conclusion")
    set_cell(table2.cell(2, 1), "To be completed by the supervisor.")


def main():
    shutil.copyfile(TEMPLATE, OUT)
    doc = Document(OUT)
    fill_doc(doc)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
